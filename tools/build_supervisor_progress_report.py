from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parents[2]
SUMMARY_PATH = ROOT / "prototype" / "output" / "carla_lead_braking" / "summary.json"
OUTPUT_PATH = ROOT / "Supervisor Progress Report - 3 September 2026.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1F2937"
MUTED = "5B6573"
LIGHT_GRAY = "F2F4F7"
LIGHT_BLUE = "EAF2F8"
LIGHT_GOLD = "FFF6D8"
LIGHT_GREEN = "EAF6EE"
WHITE = "FFFFFF"
RED = "9B1C1C"
GOLD = "7A5A00"
GREEN = "276749"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def set_font(run, *, name="Calibri", size=11, bold=None, italic=None, color=INK):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color="C8CDD4", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa):
    if sum(widths_dxa) != TABLE_WIDTH_DXA:
        raise ValueError(f"Column widths must sum to {TABLE_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Twips(widths_dxa[index])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def configure_numbering(doc):
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(el.get(qn("w:abstractNumId"))) for el in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(el.get(qn("w:numId"))) for el in numbering.findall(qn("w:num"))]
    next_abstract = max(abstract_ids, default=0) + 1
    next_num = max(num_ids, default=0) + 1

    def add_num_instance(abstract_id, *, restart=False):
        nonlocal next_num
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(next_num))
        abstract_id_element = OxmlElement("w:abstractNumId")
        abstract_id_element.set(qn("w:val"), str(abstract_id))
        num.append(abstract_id_element)
        if restart:
            level_override = OxmlElement("w:lvlOverride")
            level_override.set(qn("w:ilvl"), "0")
            start_override = OxmlElement("w:startOverride")
            start_override.set(qn("w:val"), "1")
            level_override.append(start_override)
            num.append(level_override)
        numbering.append(num)
        instance_id = next_num
        next_num += 1
        return instance_id

    def add_definition(style_name, num_format, level_text):
        nonlocal next_abstract, next_num
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(next_abstract))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        level = OxmlElement("w:lvl")
        level.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        level.append(start)
        fmt = OxmlElement("w:numFmt")
        fmt.set(qn("w:val"), num_format)
        level.append(fmt)
        text = OxmlElement("w:lvlText")
        text.set(qn("w:val"), level_text)
        level.append(text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        level.append(suff)
        ppr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        ppr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        ppr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        ppr.append(spacing)
        level.append(ppr)
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        fonts.set(qn("w:ascii"), "Calibri")
        fonts.set(qn("w:hAnsi"), "Calibri")
        rpr.append(fonts)
        level.append(rpr)
        abstract.append(level)
        first_num_index = next(
            (index for index, child in enumerate(numbering) if child.tag == qn("w:num")),
            len(numbering),
        )
        numbering.insert(first_num_index, abstract)
        abstract_id = next_abstract
        num_id = add_num_instance(abstract_id)
        style = doc.styles.add_style(style_name, 1)
        style.base_style = doc.styles["Normal"]
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167
        next_abstract += 1
        return num_id, abstract_id

    bullet_num, _ = add_definition("Meeting Bullet", "bullet", "\u2022")
    number_rq, number_abstract = add_definition("Meeting Number", "decimal", "%1.")
    return {
        "bullet": bullet_num,
        "number_rq": number_rq,
        "number_questions": add_num_instance(number_abstract, restart=True),
    }


def add_list_item(doc, text, num_id, style_name):
    paragraph = doc.add_paragraph(style=style_name)
    num_pr = paragraph._p.get_or_add_pPr().get_or_add_numPr()
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num)
    set_font(paragraph.add_run(text))
    return paragraph


def add_callout(doc, label, text, *, fill=LIGHT_BLUE, border=BLUE):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.line_spacing = 1.10
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), border)
    borders.append(left)
    p_pr.append(borders)
    set_font(paragraph.add_run(f"{label}: "), bold=True, color=DARK_BLUE)
    set_font(paragraph.add_run(text))
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(text, style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_body(doc, text, *, bold_lead=None, before=0):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    if bold_lead and text.startswith(bold_lead):
        set_font(paragraph.add_run(bold_lead), bold=True)
        set_font(paragraph.add_run(text[len(bold_lead):]))
    else:
        set_font(paragraph.add_run(text))
    return paragraph


def add_table(doc, headers, rows, widths, *, alignments=None, font_size=9.2):
    table = doc.add_table(rows=1, cols=len(headers))
    set_table_geometry(table, widths)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for index, header in enumerate(headers):
        cell = table.rows[0].cells[index]
        set_cell_shading(cell, LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.05
        set_font(paragraph.add_run(header), size=font_size, bold=True, color=DARK_BLUE)
    for row_values in rows:
        row_cells = table.add_row().cells
        for index, value in enumerate(row_values):
            paragraph = row_cells[index].paragraphs[0]
            paragraph.alignment = (alignments[index] if alignments else WD_ALIGN_PARAGRAPH.LEFT)
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            set_font(paragraph.add_run(str(value)), size=font_size, color=INK)
    set_table_geometry(table, widths)
    return table


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_font(run, size=9, color=MUTED)


def configure_document(doc):
    doc.settings.odd_and_even_pages_header_footer = True
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    heading_tokens = {
        1: (16, BLUE, 16, 8),
        2: (13, BLUE, 12, 6),
        3: (12, DARK_BLUE, 8, 4),
    }
    for level, (size, color, before, after) in heading_tokens.items():
        style = doc.styles[f"Heading {level}"]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    def populate_header(paragraph):
        paragraph.clear()
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
        set_font(paragraph.add_run("MSc Dissertation Progress Report"), size=9, bold=True, color=MUTED)
        set_font(paragraph.add_run("\tSupervisor Meeting | 3 September 2026"), size=9, color=MUTED)

    def populate_footer(paragraph):
        paragraph.clear()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(0)
        set_font(paragraph.add_run("Page "), size=9, color=MUTED)
        add_field(paragraph, " PAGE ")

    populate_header(section.header.paragraphs[0])
    populate_header(section.even_page_header.paragraphs[0])
    populate_footer(section.footer.paragraphs[0])
    populate_footer(section.even_page_footer.paragraphs[0])
    section.first_page_header.paragraphs[0].clear()
    section.first_page_footer.paragraphs[0].clear()

    doc.core_properties.title = "MSc Dissertation Progress Report"
    doc.core_properties.subject = "Supervisor meeting progress update"
    doc.core_properties.author = "Zainual Mohammed"
    doc.core_properties.keywords = "MSc dissertation, CARLA, XMAI, progress report"


def build_report():
    summary = json.loads(SUMMARY_PATH.read_text(encoding="utf-8"))
    results = {item["condition"]: item for item in summary["conditions"]}

    doc = Document()
    configure_document(doc)
    numbering = configure_numbering(doc)

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(8)
    kicker.paragraph_format.space_after = Pt(3)
    set_font(kicker.add_run("MSC DISSERTATION | PROGRESS UPDATE"), size=9.5, bold=True, color=BLUE)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    set_font(title.add_run("Supervisor Progress Report"), size=24, bold=True, color="111827")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.paragraph_format.keep_with_next = True
    set_font(
        subtitle.add_run(
            "An Explainable Multi-Agent AI Framework for Safe and Trustworthy Autonomous Vehicle Decision-Making"
        ),
        size=13,
        color=MUTED,
    )

    for label, value in (
        ("Student", "Zainual Mohammed"),
        ("Meeting date", "3 September 2026"),
        ("Submission deadline", "10 September 2026"),
        ("Current stage", "Data Preparation; initial Analysis and Results started"),
    ):
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(0)
        paragraph.paragraph_format.space_after = Pt(2)
        set_font(paragraph.add_run(f"{label}: "), bold=True, color=DARK_BLUE)
        set_font(paragraph.add_run(value))

    separator = doc.add_paragraph()
    separator.paragraph_format.space_before = Pt(8)
    separator.paragraph_format.space_after = Pt(8)
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "14")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), BLUE)
    p_bdr.append(bottom)
    separator._p.get_or_add_pPr().append(p_bdr)

    add_callout(
        doc,
        "Current position",
        "The literature review is complete. The research scope, questions, architecture, evidence format, tested prototype, CARLA installation, and first live CARLA proof-of-concept experiment are complete. The immediate priority is repeated simulator evaluation and dissertation writing.",
    )

    add_heading(doc, "1. Project aim and MSc scope")
    add_body(
        doc,
        "The project designs and evaluates a lightweight Explainable Multi-Agent AI (XMAI) pipeline for autonomous-vehicle hazard-response decisions in CARLA. The intended contribution is a transparent decision process in which each action and explanation can be reconstructed from structured evidence. The study is deliberately limited to simulation and does not claim production readiness, real-vehicle validation, or safety certification.",
    )
    add_body(doc, "The working research questions address:")
    for text in (
        "How functional agents can use structured scene evidence to produce traceable hazard-response decisions in a bounded urban-driving domain.",
        "How the prototype compares with a matched centralised baseline on collision outcome, minimum time-to-collision, response behaviour, and decision latency.",
        "Whether evidence-linked causal and counterfactual explanations improve transparency and can be verified against the decision trace.",
    ):
        add_list_item(doc, text, numbering["number_rq"], "Meeting Number")

    add_heading(doc, "2. Progress against the supervisor milestones")
    milestone_rows = [
        ("1", "Topic confirmation and kickoff", "Complete", "Working title, aim, research questions, and MSc boundary defined."),
        ("2", "Proposal development", "Draft available", "Research Proposal.pdf is retained; final MSc-scale alignment still needs review."),
        ("3", "Proposal approval and literature review start", "Partly confirmed", "Literature review complete; formal proposal approval evidence is not recorded in the workspace."),
        ("4", "Literature review and methodology design", "Complete", "Review, methodology notes, evaluation blueprint, agent interfaces, and evidence specification completed."),
        ("5", "Data preparation", "In progress", "CARLA 0.9.16 installed; prototype tested; first three-condition live run completed."),
        ("6", "Analysis and results", "Started", "Initial single-seed table and interpretation produced; repeated evaluation remains."),
        ("7", "Discussion and full draft", "Next", "Writing will proceed alongside the remaining experiments."),
        ("8", "Revision and submission", "Scheduled", "Final checking and submission planned for 9-10 September."),
    ]
    add_table(
        doc,
        ["Stage", "Milestone", "Status", "Evidence / next requirement"],
        milestone_rows,
        [620, 2590, 1300, 4850],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=8.7,
    )

    add_heading(doc, "3. Work completed")
    completed_items = (
        "Completed the literature review and retained the research proposal, research questions, and dataset access references.",
        "Defined the operational boundary: one CARLA town initially, dry daytime conditions, low-to-moderate urban speeds, one ego vehicle, and a small number of relevant actors.",
        "Designed five functional roles: Scene-State, Risk, Prediction, Decision, and Explainability.",
        "Implemented a shared action set: CONTINUE, SLOW_YIELD, and CONTROLLED_STOP.",
        "Defined a matched centralised baseline and two XMAI conditions so the same inputs, thresholds, actions, and seeds can be compared.",
        "Implemented structured JSONL evidence traces, CSV measurements, unique decision-step identifiers, causal explanations, counterfactual explanations, and automatic faithfulness validation.",
        "Completed a deterministic scripted pilot: 90 records across three scenarios, three conditions, and ten seeds; 30 explanations were generated and all 30 passed evidence-link validation.",
        "Installed and verified CARLA 0.9.16 with a project-local Python environment and live-world adapter.",
        "Completed the first live lead-vehicle-braking proof of concept in CARLA under all three experimental conditions.",
        "Verified the current code with nine automated tests, all passing.",
    )
    for text in completed_items:
        add_list_item(doc, text, numbering["bullet"], "Meeting Bullet")

    add_heading(doc, "4. Current prototype and experiment design")
    add_body(
        doc,
        "CARLA supplies ground-truth simulator state to a small SceneState interface. The Risk and Prediction roles assess the hazard and time to conflict; the Decision role evaluates permitted actions; the Explainability role then produces causal and counterfactual statements using only evidence that already exists in the decision trace. Explanation generation remains outside the vehicle-control decision path.",
    )
    for text in (
        "Centralised baseline: matched policy with a minimal decision record.",
        "XMAI decision: Scene-State, Risk, Prediction, and Decision roles with a structured trace.",
        "XMAI with explanation: the same decision pipeline plus evidence-linked causal and counterfactual explanations.",
    ):
        add_list_item(doc, text, numbering["bullet"], "Meeting Bullet")

    add_heading(doc, "5. Initial live CARLA result")
    add_body(
        doc,
        "The lead-vehicle-braking proof of concept was run in Carla/Maps/Town01 with seed 0, identical recorded spawn transforms, synchronous mode, and a fixed 0.05-second simulation interval. The following are observed simulator values, not hypothetical results.",
    )

    baseline = results["centralised_baseline"]
    decision = results["xmai_decision"]
    explanation = results["xmai_explanation"]
    result_rows = []
    for label, item, validity in (
        ("Centralised baseline", baseline, "Not generated"),
        ("XMAI decision", decision, "Not generated"),
        ("XMAI + explanation", explanation, f'{item_value(explanation, "valid_explanations")}/{item_value(explanation, "generated_explanations")}'),
    ):
        result_rows.append(
            (
                label,
                item["ticks"],
                item["collision_count"],
                f'{item["minimum_clearance_m"]:.4f}',
                f'{item["minimum_ttc_s"]:.4f}',
                f'{item["mean_decision_latency_ms"]:.6f}',
                validity,
            )
        )
    add_table(
        doc,
        ["Condition", "Ticks", "Crashes", "Min clearance (m)", "Min TTC (s)", "Mean latency (ms)", "Valid explanations"],
        result_rows,
        [1980, 680, 850, 1290, 1050, 1420, 2090],
        alignments=[WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 6,
        font_size=8.2,
    )

    decision_overhead = decision["mean_decision_latency_ms"] - baseline["mean_decision_latency_ms"]
    explanation_overhead = explanation["mean_decision_latency_ms"] - decision["mean_decision_latency_ms"]
    add_body(
        doc,
        f"All three retained conditions followed the same 35-tick action sequence, recorded identical clearance and TTC outcomes, and stopped without collision. The structured decision path added {decision_overhead:.6f} ms mean processing time relative to the minimal baseline in this run. The explanation-enabled condition produced {explanation['generated_explanations']} explanations; all {explanation['valid_explanations']} passed the evidence-link validation rule. Its measured mean processing time was {explanation_overhead:.6f} ms above the decision-only pipeline.",
        before=6,
    )
    add_callout(
        doc,
        "Important limitation",
        "This is one seed and one scenario, so it demonstrates that the end-to-end system works but does not support statistical or general safety claims. Repeated fixed-seed runs, initial-state equivalence checks, and transparent reporting of failed or variable runs are required before comparing the conditions in the dissertation.",
        fill=LIGHT_GOLD,
        border=GOLD,
    )

    add_heading(doc, "6. What I am doing now")
    for text in (
        "Strengthening the lead-braking experiment with pre-run checks for vehicle position, speed, and clearance equivalence.",
        "Preparing repeated fixed-seed runs for the three experimental conditions.",
        "Retaining per-tick CSV data, full JSONL evidence traces, condition summaries, configuration files, and failed-run information for reproducibility.",
        "Starting the methodology and implementation write-up while experimental work continues.",
        "Keeping the project at MSc scale so the final submission is coherent and evidence based by 10 September.",
    ):
        add_list_item(doc, text, numbering["bullet"], "Meeting Bullet")

    add_heading(doc, "7. Plan to submission")
    plan_rows = [
        ("3-4 September", "Data preparation", "Stabilise and repeat lead-braking runs; add pedestrian and junction scenarios only after the first scenario is reliable."),
        ("5-6 September", "Analysis and results", "Export analysis-ready tables and figures; report safety, latency, trace completeness, explanation validity, failures, and uncertainty."),
        ("7-8 September", "Discussion and full draft", "Integrate methodology, implementation, results, limitations, discussion, conclusion, and references."),
        ("9-10 September", "Revision and submission", "Proofread, verify citations and formatting, inspect the final PDF, submit early, and retain the receipt."),
    ]
    add_table(
        doc,
        ["Dates", "Stage", "Planned output"],
        plan_rows,
        [1800, 2180, 5380],
        alignments=[WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        font_size=9,
    )

    add_heading(doc, "8. Decisions requested from the supervisor")
    supervisor_questions = (
        "Confirm the official word count, formatting rules, submission time, required filename, and whether references or appendices are excluded from the limit.",
        "Confirm that the narrowed MSc contribution and three research questions are appropriate for submission.",
        "Confirm whether one robust scenario with repeated seeds is preferable to three less-developed scenarios if time becomes constrained.",
        "Advise whether a participant-based explanation study is required. With the current deadline, the defensible fallback is automatic faithfulness validation and a clearly labelled exploratory assessment only if prior ethics approval exists.",
        "Confirm the proposal approval status and identify any mandatory corrections that must be reflected in the final dissertation.",
        "Agree the priority order for the remaining week: reproducible evidence first, then complete write-up, then optional extensions.",
    )
    for text in supervisor_questions:
        add_list_item(doc, text, numbering["number_questions"], "Meeting Number")

    add_heading(doc, "9. Evidence available for demonstration")
    for text in (
        "CARLA 0.9.16 running locally with connection-verification scripts.",
        "Lead-braking experiment runner and fixed configuration files.",
        "Per-tick CSV results, condition summaries, and complete JSONL evidence traces.",
        "An example causal and counterfactual explanation linked to the exact decision, risk, and prediction evidence IDs.",
        "Nine passing automated tests and the 90-record scripted pilot output.",
        "Project plan, methodology notes, evaluation blueprint, agent-interface specification, and proof-of-concept results note.",
    ):
        add_list_item(doc, text, numbering["bullet"], "Meeting Bullet")

    add_heading(doc, "10. Suggested one-minute verbal update")
    add_callout(
        doc,
        "Meeting summary",
        f"My literature review is complete and I have narrowed the project to an MSc-scale CARLA proof of concept. I designed and tested a five-role explainable decision pipeline, a matched centralised baseline, structured evidence logging, and automatic explanation validation. CARLA 0.9.16 is installed and the first lead-vehicle-braking experiment has run in all three conditions. All three observed runs avoided collision, and all {explanation['generated_explanations']} generated explanations passed validation. This is still a single-seed proof of concept, so I am now strengthening reproducibility and preparing repeated runs before making comparative claims. I need your confirmation on the scope, word count, proposal approval, and whether repeated depth in one scenario should take priority over breadth across three scenarios before the 10 September deadline.",
        fill=LIGHT_GREEN,
        border=GREEN,
    )

    doc.save(OUTPUT_PATH)
    return OUTPUT_PATH


def item_value(item, key):
    return item[key]


if __name__ == "__main__":
    print(build_report())
